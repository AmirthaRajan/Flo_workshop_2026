"""Load workshop knowledge sources into LangChain Documents."""

import email
import email.policy
import ipaddress
import os
import socket
import zipfile
from pathlib import Path
from urllib.parse import urlparse

from bs4 import BeautifulSoup
from docx import Document as WordDocument
from langchain_core.documents import Document
from pypdf import PdfReader

from config import settings

DOCUMENTS_DIR = Path("data/documents")


def save_uploaded_file(uploaded_file) -> Path:
    """Save an upload locally so it can be inspected after the workshop."""
    DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)
    # Keep only the filename so an upload cannot write outside the data folder.
    file_path = DOCUMENTS_DIR / Path(uploaded_file.name).name
    file_path.write_bytes(uploaded_file.getvalue())
    return file_path


def load_file(file_path: str | Path) -> list[Document]:
    """Load a PDF or DOCX file into a shared document representation."""
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        reader = PdfReader(path)
        documents = []
        for page_number, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                documents.append(
                    Document(
                        page_content=text,
                        metadata={"source": path.name, "page": page_number + 1},
                    )
                )
        return documents

    if suffix == ".docx":
        if not zipfile.is_zipfile(path):
            # Confluence "Export to Word" produces an MHTML file (MIME-wrapped
            # HTML) with a .docx extension rather than a real Word document.
            return [_load_mhtml(path)]
        word_document = WordDocument(path)
        text = "\n".join(
            paragraph.text
            for paragraph in word_document.paragraphs
            if paragraph.text.strip()
        )
        return [Document(page_content=text, metadata={"source": path.name})]

    raise ValueError("Unsupported file type. Please upload a PDF or DOCX file.")


def _load_mhtml(path: Path) -> Document:
    """Extract readable text from an MHTML file such as a Confluence Word export."""
    message = email.message_from_bytes(path.read_bytes(), policy=email.policy.default)
    html_part = next(
        (part for part in message.walk() if part.get_content_type() == "text/html"),
        None,
    )
    if html_part is None:
        raise ValueError(
            f"{path.name} is not a valid Word document. Re-save it as .docx in "
            "Word or export the page as PDF."
        )
    soup = BeautifulSoup(html_part.get_content(), "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    text = "\n".join(
        line.strip() for line in soup.get_text("\n").splitlines() if line.strip()
    )
    if not text:
        raise ValueError(f"No readable text was found in {path.name}.")
    return Document(page_content=text, metadata={"source": path.name})


def load_url(url: str) -> list[Document]:
    """Extract readable text from a public or locally accessible web page."""
    # WebBaseLoader reads this setting when imported, so set it first.
    os.environ.setdefault("USER_AGENT", "RAG-Workshop/1.0")
    from langchain_community.document_loaders import WebBaseLoader

    parsed_url = urlparse(url)
    if parsed_url.scheme not in {"http", "https"} or not parsed_url.netloc:
        raise ValueError("Enter a complete URL beginning with http:// or https://.")
    if parsed_url.username or parsed_url.password:
        raise ValueError("URLs containing usernames or passwords are not supported.")

    if not settings.allow_private_urls:
        try:
            addresses = {
                result[4][0]
                for result in socket.getaddrinfo(parsed_url.hostname, None)
            }
        except socket.gaierror as error:
            raise ValueError("The URL hostname could not be resolved.") from error
        if not addresses or any(
            not ipaddress.ip_address(address).is_global for address in addresses
        ):
            raise ValueError(
                "Private network URLs are blocked. Set ALLOW_PRIVATE_URLS=true "
                "only when intentionally loading a trusted internal wiki."
            )

    web_loader = WebBaseLoader(
        web_paths=[url],
        header_template={"User-Agent": "RAG-Workshop/1.0"},
        requests_kwargs={"timeout": 15, "allow_redirects": False},
        raise_for_status=True,
        show_progress=False,
    )
    loaded_documents = web_loader.load()
    text = "\n".join(
        document.page_content.strip()
        for document in loaded_documents
        if document.page_content.strip()
    )
    if not text:
        raise ValueError("No readable text was found at that URL.")
    return [Document(page_content=text, metadata={"source": url})]
