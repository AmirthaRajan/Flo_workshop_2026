"""Focused tests for provider-independent workshop logic."""

import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from langchain_core.documents import Document

import loaders
import rag
from config import Settings


class ConfigTests(unittest.TestCase):
    def test_embedding_and_model_providers_can_differ(self):
        with patch.dict(
            os.environ,
            {"MODEL_PROVIDER": "openai", "EMBEDDING_PROVIDER": "ollama"},
        ):
            configured = Settings(
                model_provider=os.environ["MODEL_PROVIDER"],
                embedding_provider=os.environ["EMBEDDING_PROVIDER"],
            )

        self.assertEqual(configured.model_provider, "openai")
        self.assertEqual(configured.embedding_provider, "ollama")

    def test_overlap_must_be_smaller_than_chunk_size(self):
        with self.assertRaisesRegex(ValueError, "CHUNK_OVERLAP"):
            Settings(chunk_size=100, chunk_overlap=100)


class LoaderTests(unittest.TestCase):
    def test_docx_loader_preserves_source_name(self):
        from docx import Document as WordDocument

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "guide.docx"
            document = WordDocument()
            document.add_paragraph("Run the project locally.")
            document.save(path)

            loaded = loaders.load_file(path)

        self.assertEqual(loaded[0].page_content, "Run the project locally.")
        self.assertEqual(loaded[0].metadata["source"], "guide.docx")

    def test_url_requires_http_or_https(self):
        with self.assertRaisesRegex(ValueError, "complete URL"):
            loaders.load_url("example.com/docs")


class RagTests(unittest.TestCase):
    def test_split_documents_keeps_source_metadata(self):
        documents = [
            Document(page_content="project setup " * 100, metadata={"source": "guide"})
        ]

        chunks = rag.split_documents(documents)

        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(chunk.metadata["source"] == "guide" for chunk in chunks))

    def test_prompt_contains_context_question_and_fallback(self):
        prompt = rag.build_prompt("Use Python 3.11.", "Which Python version?")

        self.assertIn("Use Python 3.11.", prompt)
        self.assertIn("Which Python version?", prompt)
        self.assertIn(rag.FALLBACK_ANSWER, prompt)


if __name__ == "__main__":
    unittest.main()
