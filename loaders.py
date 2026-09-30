"""Load workshop knowledge sources into LangChain Documents."""

import ipaddress
import socket
from pathlib import Path
from urllib.parse import urlparse

import requests
from bs4 import BeautifulSoup
from docx import Document as WordDocument
from langchain_core.documents import Document
from pypdf import PdfReader

from config import settings

DOCUMENTS_DIR = Path("data/documents")


def save_uploaded_file(uploaded_file) -> Path:
    """Save an upload locally so it can be inspected after the workshop."""
    DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)
    # Keep only the filename so an upload cannot write outside the data folder.
    file_path = DOCUMENTS_DIR / Path(uploaded_file.name).name
    file_path.write_bytes(uploaded_file.getvalue())
    return file_path


def load_file(file_path: str | Path) -> list[Document]:
    """Load a PDF or DOCX file into a shared document representation."""
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        reader = PdfReader(path)
        documents = []
        for page_number, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                documents.append(
                    Document(
                        page_content=text,
                        metadata={"source": path.name, "page": page_number + 1},
                    )
                )
        return documents

    if suffix == ".docx":
        word_document = WordDocument(path)
        text = "\n".join(
            paragraph.text
            for paragraph in word_document.paragraphs
            if paragraph.text.strip()
        )
        return [Document(page_content=text, metadata={"source": path.name})]

    raise ValueError("Unsupported file type. Please upload a PDF or DOCX file.")


def load_url(url: str) -> list[Document]:
    """Extract readable text from a public or locally accessible web page."""
    parsed_url = urlparse(url)
    if parsed_url.scheme not in {"http", "https"} or not parsed_url.netloc:
        raise ValueError("Enter a complete URL beginning with http:// or https://.")
    if parsed_url.username or parsed_url.password:
        raise ValueError("URLs containing usernames or passwords are not supported.")

    if not settings.allow_private_urls:
        try:
            addresses = {
                result[4][0]
                for result in socket.getaddrinfo(parsed_url.hostname, None)
            }
        except socket.gaierror as error:
            raise ValueError("The URL hostname could not be resolved.") from error
        if not addresses or any(
            not ipaddress.ip_address(address).is_global for address in addresses
        ):
            raise ValueError(
                "Private network URLs are blocked. Set ALLOW_PRIVATE_URLS=true "
                "only when intentionally loading a trusted internal wiki."
            )

    response = requests.get(
        url,
        headers={"User-Agent": "RAG-Workshop/1.0"},
        timeout=15,
        allow_redirects=False,
    )
    if response.is_redirect:
        raise ValueError("Redirecting URLs are not supported; enter the final URL.")
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "html.parser")
    for unwanted in soup(["script", "style", "nav", "footer"]):
        unwanted.decompose()

    text = "\n".join(soup.stripped_strings)
    if not text:
        raise ValueError("No readable text was found at that URL.")
    return [Document(page_content=text, metadata={"source": url})]
